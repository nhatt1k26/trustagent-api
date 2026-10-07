"""Tiện ích xuất file Excel (.xlsx) dùng openpyxl.

Tách riêng để tái sử dụng cho nhiều loại báo cáo (đội ngũ, thù lao...).
"""

from datetime import datetime
from io import BytesIO
from typing import Any, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Style dùng chung
_HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TITLE_FONT = Font(bold=True, size=14)
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build_sheet_xlsx(
    title: str,
    headers: Sequence[str],
    rows: List[Sequence[Any]],
    sheet_name: str = "Sheet1",
    subtitle: str | None = None,
) -> bytes:
    """Tạo 1 file xlsx đơn giản có tiêu đề + bảng dữ liệu. Trả về bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31] or "Sheet1"

    ncols = len(headers)
    last_col = get_column_letter(ncols)

    # Dòng tiêu đề (merge toàn bộ chiều rộng)
    ws.merge_cells(f"A1:{last_col}1")
    tcell = ws["A1"]
    tcell.value = title
    tcell.font = _TITLE_FONT
    tcell.alignment = _CENTER

    header_row = 2
    if subtitle:
        ws.merge_cells(f"A2:{last_col}2")
        scell = ws["A2"]
        scell.value = subtitle
        scell.alignment = _CENTER
        scell.font = Font(italic=True, size=10, color="6B7280")
        header_row = 3

    # Header
    for col_idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER

    # Dữ liệu
    for r_off, row in enumerate(rows, start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_off, column=c_idx, value=value)
            cell.border = _BORDER
            cell.alignment = _LEFT if isinstance(value, str) else _CENTER

    # Tự canh độ rộng cột theo nội dung
    for c_idx in range(1, ncols + 1):
        letter = get_column_letter(c_idx)
        max_len = len(str(headers[c_idx - 1]))
        for row in rows:
            val = row[c_idx - 1] if c_idx - 1 < len(row) else ""
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)

    # Freeze header
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def timestamped_filename(prefix: str) -> str:
    """Sinh tên file kèm timestamp, vd: doi_ngu_20260930_1530.xlsx"""
    return f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
